import io
from zipfile import ZipFile
import pytest
from openpyxl import load_workbook
from peanalyzer.gui.sample_workbook import template,parse,COLUMNS
from peanalyzer.barcodes import load_catalog,bundled_catalog


def workbook(rows=None,mutate=None):
    book=load_workbook(io.BytesIO(template()));sheet=book['Samples']
    for row in rows or [dict(barcode_id='4A',sample_name='leaf1',editor='PE',treatment='test',dpi=6,biological_replicate='1')]:
        sheet.append([row.get(k) for k in COLUMNS])
    if mutate:mutate(book,sheet)
    stream=io.BytesIO();book.save(stream);return stream.getvalue()


def test_template_and_catalog_resolution():
    result=parse(workbook(),'qi_hitom_v1');assert result['errors']==[]
    row=result['samples'][0];catalog=load_catalog(bundled_catalog())
    assert row['forward_barcode']==catalog.ends['forward']['F4']
    assert row['reverse_barcode']==catalog.ends['reverse']['RA']
    assert row['dpi']=='6' and row['paired'] is False and row['block_id']==''
    assert result['imported_count']==1


def test_twenty_samples():
    rows=[dict(barcode_id=f'{i//8+1}{chr(65+i%8)}',sample_name=f's{i}',editor='PE',treatment='t',dpi=6,biological_replicate=str(i+1)) for i in range(20)]
    result=parse(workbook(rows),'qi_hitom_v1');assert not result['errors'],result
    assert result['imported_count']==20


@pytest.mark.parametrize('field,value',[('barcode_id','9Z'),('sample_name',''),('editor',''),('biological_replicate',''),('dpi','six'),('dpi',1.5),('dpi',-1),('paired','maybe'),('paired',True),('forward_barcode','NNNN'),('forward_barcode_id','F5')])
def test_row_field_errors(field,value):
    def change(book,sheet):sheet.cell(2,COLUMNS.index(field)+1,value)
    result=parse(workbook(mutate=change),'qi_hitom_v1')
    assert result['samples']==[] and result['errors']
    assert any(e['row']==2 for e in result['errors'])


@pytest.mark.parametrize('attack',['missing_column','unknown_column','duplicate_column','duplicate_barcode','duplicate_sample','blank_row','formula','extra_sheet','hidden_sheet','hidden_row','hidden_column','merged','named_range','bad_version','hyperlink','limit'])
def test_workbook_structure(attack):
    def change(book,sheet):
        if attack=='missing_column':sheet.cell(1,5,'block_id')
        elif attack=='unknown_column':sheet.cell(1,1,'barcod_id')
        elif attack=='duplicate_column':sheet.cell(1,2,'barcode_id')
        elif attack.startswith('duplicate_'):
            for c in range(1,len(COLUMNS)+1):sheet.cell(3,c,sheet.cell(2,c).value)
            if attack=='duplicate_sample':sheet.cell(3,1,'5A')
        elif attack=='blank_row':sheet.cell(4,1,'5A')
        elif attack=='formula':sheet.cell(2,5,'=3+3')
        elif attack=='extra_sheet':book.create_sheet('Other')
        elif attack=='hidden_sheet':sheet.sheet_state='hidden'
        elif attack=='hidden_row':sheet.row_dimensions[2].hidden=True
        elif attack=='hidden_column':sheet.column_dimensions['A'].hidden=True
        elif attack=='merged':sheet.merge_cells('A2:B2')
        elif attack=='named_range':
            from openpyxl.workbook.defined_name import DefinedName
            book.defined_names.add(DefinedName('hidden',attr_text="Samples!$A$2"))
        elif attack=='bad_version':book['Instructions']['A1']='other'
        elif attack=='hyperlink':sheet.cell(2,2).hyperlink='https://example.invalid'
        elif attack=='limit':sheet.cell(1002,1,'4A')
    result=parse(workbook(mutate=change),'qi_hitom_v1')
    assert result['errors'] and result['samples']==[]


@pytest.mark.parametrize('paired',[None,'false',' FALSE ',False])
def test_safe_blank_and_whitespace(paired):
    rows=[dict(barcode_id=' 4A ',sample_name=' leaf1 ',editor=' PE ',treatment=' t ',dpi=' 6 ',biological_replicate=' 01 ',paired=paired)]
    row=parse(workbook(rows),'qi_hitom_v1')['samples'][0]
    assert row['sample_name']=='leaf1' and row['biological_replicate']=='01' and row['paired'] is False


def test_custom_and_boolean_true():
    rows=[dict(barcode_id='s1',sample_name='leaf1',editor='PE',treatment='t',dpi=6,biological_replicate='1',paired=True,block_id='leaf',forward_barcode='ACGT',reverse_barcode='TGCA')]
    result=parse(workbook(rows),'custom');assert not result['errors'],result
    assert result['samples'][0]['paired'] is True


def test_zip_external_and_invalid_content():
    assert parse(b'not an xlsx','qi_hitom_v1')['errors']
    source=workbook();out=io.BytesIO()
    with ZipFile(io.BytesIO(source)) as inp,ZipFile(out,'w') as dest:
        for entry in inp.infolist():dest.writestr(entry,inp.read(entry.filename))
        dest.writestr('xl/vbaProject.bin',b'macro')
    assert parse(out.getvalue(),'qi_hitom_v1')['errors']
