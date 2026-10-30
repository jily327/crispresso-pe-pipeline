"""Strict XLSX transport adapter to existing Sample/catalog/CSV authorities."""
import io
import re
from zipfile import ZipFile, BadZipFile

from defusedxml.ElementTree import fromstring
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from peanalyzer.barcodes import Catalog, bundled_catalog, load_catalog
from peanalyzer.schemas import Sample
from peanalyzer.validation import _csv_value, load_samples
from peanalyzer.gui.setup_validation import MemoryCSV
from peanalyzer.gui.setup_models import DraftSample
from peanalyzer.gui.workspace import ViewerError

VERSION='PEAnalyzer sample template v1'
COLUMNS=('barcode_id','sample_name','editor','treatment','dpi','biological_replicate','block_id','paired',
         'forward_barcode_id','reverse_barcode_id','forward_barcode','reverse_barcode')
REQUIRED=set(COLUMNS[:6])
MAX_BYTES=2*1024**2
MAX_ROWS=1000


def template():
    book=Workbook();sheet=book.active;sheet.title='Samples';sheet.append(COLUMNS)
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width=24
    sheet.freeze_panes='A2'
    info=book.create_sheet('Instructions',0)
    for line in [VERSION,'One biological sample per row; do not insert blank rows within data.',
                 'Required: barcode_id, sample_name, editor, treatment, dpi, biological_replicate.',
                 'QI-HITOM v1: e.g. 4A resolves F4/RA through the qualified catalog; no sequences needed.',
                 'Custom profile: supply forward_barcode and reverse_barcode sequences.',
                 'Optional paired: blank/false means unpaired; true requires block_id.',
                 'No formulas, macros, hidden rows/sheets, merged cells, tables, links or extra sheets.',
                 'Use text cells for IDs to preserve leading zeros. Maximum 1000 data rows.',
                 'Import populates editable Samples only. Existing Validate & Review is still required.']:
        info.append([line])
    info.column_dimensions['A'].width=110
    output=io.BytesIO();book.save(output);return output.getvalue()


def issue(row,field,message,code='WORKBOOK_VALIDATION'):
    return dict(row=row,field=field,code=code,severity='ERROR',message=message)


def inspect_zip(data):
    if len(data)>MAX_BYTES:raise ValueError('Workbook exceeds 2 MiB.')
    with ZipFile(io.BytesIO(data)) as archive:
        entries=archive.infolist();names=[e.filename for e in entries]
        if len(entries)>100 or len(set(names))!=len(names) or sum(e.file_size for e in entries)>16*1024**2:
            raise ValueError('Workbook archive exceeds the supported structure/expanded-size limit.')
        if any(e.flag_bits & 1 for e in entries):raise ValueError('Encrypted workbooks are unsupported.')
        allowed={'[Content_Types].xml','_rels/.rels','docProps/app.xml','docProps/core.xml',
                 'xl/styles.xml','xl/sharedStrings.xml','xl/theme/theme1.xml','xl/workbook.xml',
                 'xl/_rels/workbook.xml.rels'}
        for name in names:
            if name not in allowed and not re.fullmatch(r'xl/worksheets/(?:sheet[12]\.xml|_rels/sheet[12]\.xml\.rels)',name):
                raise ValueError('Unsupported workbook part: '+name)
            if name.startswith('/') or '\\' in name or '..' in name.split('/'):
                raise ValueError('Unsafe workbook archive path.')
            if any(part in name.lower() for part in ('vbaproject','externallinks/','embeddings/','customui/','connections.xml','pivot','activex/')):
                raise ValueError('Macros, embedded/external content and pivot structures are unsupported.')
            if name.endswith('.xml') or name.endswith('.rels'):
                tree=fromstring(archive.read(name))
                if name.endswith('.rels') and any(n.attrib.get('TargetMode')=='External' for n in tree.iter()):
                    raise ValueError('External workbook relationships are unsupported.')
                if name=='[Content_Types].xml' and any('macroenabled' in str(n.attrib).lower() for n in tree.iter()):
                    raise ValueError('Macro-enabled workbooks are unsupported.')


def value(cell, row, field, errors):
    if cell.data_type in ('f','e') or cell.hyperlink:
        errors.append(issue(row,field,'Formulas, error cells and hyperlinks are unsupported.'));return ''
    v=cell.value
    if v is None:return ''
    if isinstance(v,str):return v.strip()
    if isinstance(v,bool):
        if field=='paired':return 'true' if v else 'false'
    elif isinstance(v,(int,float)) and v==int(v):return str(int(v))
    errors.append(issue(row,field,'Use text or an exact integer; dates, fractions and other cell types are unsupported.'))
    return ''


def parse(data, profile):
    errors=[];rows=[]
    try:
        inspect_zip(data)
        book=load_workbook(io.BytesIO(data),data_only=False,keep_links=False)
        if set(book.sheetnames) not in ({'Samples'},{'Instructions','Samples'}) or book.defined_names:
            raise ValueError('Use Samples and optional official Instructions only; named ranges are unsupported.')
        if 'Instructions' in book and book['Instructions']['A1'].value!=VERSION:
            raise ValueError('Unsupported template version or Instructions sheet.')
        for sheet in book:
            if sheet.sheet_state!='visible' or sheet.merged_cells or sheet.tables or sheet._charts or sheet._images:
                raise ValueError('Hidden sheets, merged cells, tables, charts and images are unsupported.')
            if any(d.hidden for d in sheet.row_dimensions.values()) or any(d.hidden for d in sheet.column_dimensions.values()):
                raise ValueError('Hidden rows/columns are unsupported.')
            if sheet.max_row>MAX_ROWS+1 or sheet.max_column>len(COLUMNS):
                raise ValueError('Workbook exceeds 1000 sample rows or allowed column bounds.')
            if sheet.title=='Instructions':
                for row in sheet:
                    for cell in row:
                        if cell.data_type in ('f','e') or cell.hyperlink:raise ValueError('Instructions must contain plain informational text.')
        sheet=book['Samples']
        headers=[value(c,1,'header',errors) for c in sheet[1]]
        if len(headers)!=len(set(headers)):raise ValueError('Duplicate/blank column headers are unsupported.')
        unknown=set(headers)-set(COLUMNS);missing=REQUIRED-set(headers)
        if unknown:raise ValueError('Unknown column(s): '+', '.join(sorted(unknown)))
        if missing:raise ValueError('Missing required column(s): '+', '.join(sorted(missing)))
        data_rows=list(sheet.iter_rows(min_row=2))
        # Trailing empty formatting is harmless; holes inside an active table are not guessed away.
        while data_rows and all(c.value is None for c in data_rows[-1]):data_rows.pop()
        catalog=load_catalog(bundled_catalog()) if profile=='qi_hitom_v1' else Catalog([],authoritative=False)
        for number,cells in enumerate(data_rows,2):
            if all(c.value is None or isinstance(c.value,str) and not c.value.strip() for c in cells):
                errors.append(issue(number,'row','Blank row inside active sample table.'));continue
            raw={k:value(c,number,k,errors) for k,c in zip(headers,cells)}
            raw={k:raw.get(k,'') for k in COLUMNS}
            raw['paired']=raw['paired'] or 'false'
            if profile=='qi_hitom_v1':
                match=re.fullmatch(r'([1-9][0-9]*)([A-Z]+)',raw['barcode_id'])
                f,r=('F'+match[1],'R'+match[2]) if match else ('','')
                if f not in catalog.ends['forward'] or r not in catalog.ends['reverse']:
                    errors.append(issue(number,'barcode_id',f"{raw['barcode_id']} is not present in the qualified QI-HITOM v1 barcode catalog."));continue
                for direction,identifier in [('forward',f),('reverse',r)]:
                    field=direction+'_barcode_id'
                    if raw[field] and raw[field]!=identifier:errors.append(issue(number,field,'Catalog ID conflicts with barcode_id shorthand.'))
                    raw[field]=identifier
            try:
                converted={}
                for k,v in raw.items():
                    try:converted[k]=_csv_value(k,v)
                    except ValueError as exc:errors.append(issue(number,k,str(exc)))
                if len(converted)!=len(raw):continue
                sample=catalog.resolve_sample(Sample.model_validate(converted))
                rows.append((number,DraftSample(**{k:str(getattr(sample,k)) if k=='dpi' else getattr(sample,k) or '' for k in DraftSample.model_fields if k!='paired'},paired=sample.paired)))
            except ValidationError as exc:
                for item in exc.errors(include_url=False,include_input=False):
                    errors.append(issue(number,'.'.join(map(str,item['loc'])) or 'sample',item['msg']))
            except ValueError as exc:errors.append(issue(number,'barcode_id',str(exc)))
        if not data_rows:errors.append(issue(2,'row','Add at least one biological sample.'))
        # Existing cross-row loader is authoritative. Add precise row diagnostics for duplicate cells.
        for field in ('barcode_id','sample_name'):
            seen={}
            for number,row in rows:
                key=getattr(row,field)
                if key in seen:errors.append(issue(number,field,f'Duplicate {field}; first appears on row {seen[key]}.'))
                else:seen[key]=number
        if rows and not errors:
            try:load_samples(MemoryCSV([r.model_dump() for _,r in rows],'Excel Samples'),catalog)
            except ValueError as exc:errors.append(issue(0,'samples',str(exc)))
    except Exception as exc:
        # This adapter does not run science. Unsupported ZIP/XML/reader structures fail as import diagnostics.
        errors.append(issue(0,'workbook',str(exc)))
    finally:
        if 'book' in locals():book.close()
    return dict(template_version='samples-xlsx-v1',samples=[] if errors else [r.model_dump() for _,r in rows],
                imported_count=0 if errors else len(rows),valid_count=len(rows) if not errors else 0,
                warnings=[],errors=errors)


async def receive_workbook(request):
    if request.headers.get('content-type')!='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
        raise ViewerError('UNSUPPORTED_WORKBOOK','Expected the official .xlsx workbook body.')
    data=bytearray()
    async for block in request.stream():
        if len(data)+len(block)>MAX_BYTES:raise ViewerError('WORKBOOK_LIMIT','Workbook exceeds 2 MiB.')
        data.extend(block)
    return bytes(data)
