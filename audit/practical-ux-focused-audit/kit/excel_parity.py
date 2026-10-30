"""E: manual Sample entry vs Excel import, same registrations/targets/output -> compare canonical compiled
science/files, full preview and issues. Also: import never executes; edits after import invalidate Review."""
import json, os, shutil, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from auditor_candidate_binding import Server, Adapter, FIX, sha
S = Path(__file__).resolve().parents[2]
ORACLE = S / 'm11c/prep/m11c-audit-prep'
OUT = Path(sys.argv[1]); OUT.mkdir(parents=True, exist_ok=False)
fx = ORACLE / 'fixtures' / 'pe'; inputs = OUT / 'inputs'; inputs.mkdir()
for m in ('raw_R1.fastq.gz', 'raw_R2.fastq.gz'): shutil.copy2(fx / m, inputs / m)
outroot = OUT / 'outputs'; outroot.mkdir()
srv = Server(OUT, legacy_files=[inputs / 'raw_R1.fastq.gz', inputs / 'raw_R2.fastq.gz'], output_roots=[outroot])
refs = srv.app.state.drafts.references
ids = {p.name: k for k, p in refs.paths.items() if p.parent == inputs}; oroot = next(k for k, e in refs.entries.items() if e.kind == 'output_root')
tmpl = json.loads((fx / 'draft-template.json').read_text())
R = {}
def draft(samples, profile='qi_hitom_v1'):
    c = dict(tmpl); c.update(r1=ids['raw_R1.fastq.gz'], r2=ids['raw_R2.fastq.gz'], output_root=oroot, workspace_name='ws', samples=samples, library_profile=profile)
    d = srv.json('POST', '/api/v1/drafts').body
    d = srv.json('PATCH', f"/api/v1/drafts/{d['draft_id']}", dict(revision=d['revision'], content=c)).body
    v = srv.app.state.drafts.validate(d['draft_id'], d['revision'])       # private: same call the API makes
    return d, v
def manual_rows(case):
    from openpyxl import load_workbook
    from peanalyzer.barcodes import load_catalog, bundled_catalog
    from peanalyzer.schemas import Sample
    b = load_workbook(FIX / (case + '.xlsx')); rows = list(b['Samples'].iter_rows(values_only=True)); hdr = [h.strip() for h in rows[0]]; out = []
    for r in rows[1:]:
        v = {k: ('' if x is None else str(x).strip()) for k, x in zip(hdr, r)}
        if not v['forward_barcode']:
            bid = v['barcode_id']; s = load_catalog(bundled_catalog()).resolve_sample(Sample(sample_name=v['sample_name'], barcode_id=bid, forward_barcode_id='F' + bid[:-1],
                  reverse_barcode_id='R' + bid[-1], editor=v['editor'], treatment=v['treatment'], dpi=int(v['dpi']), biological_replicate=v['biological_replicate'], paired=False))
            v['forward_barcode'], v['reverse_barcode'] = s.forward_barcode, s.reverse_barcode
        out.append(dict(sample_name=v['sample_name'], barcode_id=v['barcode_id'], forward_barcode=v['forward_barcode'], reverse_barcode=v['reverse_barcode'], editor=v['editor'],
                        treatment=v['treatment'], dpi=v['dpi'], biological_replicate=v['biological_replicate'], paired=v['paired'].lower() == 'true', block_id=v['block_id']))
    return out
jobs_before = len(srv.app.state.jobs.jobs)
for case in ('valid-one', 'valid-twenty', 'catalog-4A', 'whitespace', 'optional-blanks'):
    r = srv.request('POST', '/api/v1/setup/samples/import?profile=qi_hitom_v1', [(FIX / (case + '.xlsx')).read_bytes()],
                    headers={'Content-Type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}).body
    assert not r['errors'], (case, r['errors'])
    dm, vm = draft(manual_rows(case)); dx, vx = draft(r['samples'])
    pm, px = vm.preview.model_dump(), vx.preview.model_dump()
    R[case] = dict(manual_status=vm.status, excel_status=vx.status, compiled_equal=vm._compiled == vx._compiled,
                   compiled_files_equal={k: vm._compiled['files'][k] == vx._compiled['files'][k] for k in vm._compiled['files']},
                   preview_equal=pm == px, preview_diff=sorted(k for k in pm if pm[k] != px[k]),
                   issues_equal=[i.model_dump() for i in vm.issues] == [i.model_dump() for i in vx.issues], n=len(r['samples']),
                   scientific_identity=(pm['scientific_identity_sha256'] == px['scientific_identity_sha256']), plan_identity=(pm['plan_identity_sha256'] == px['plan_identity_sha256']),
                   replicate_text=[s['biological_replicate'] for s in r['samples']][:3])
    print(case, json.dumps(R[case])[:500], flush=True)
# import never executes; no workspace written
R['no_execution'] = dict(jobs_created=len(srv.app.state.jobs.jobs) - jobs_before, outputs_written=sorted(os.listdir(outroot)))
# Edit after import invalidates prior Review; revalidation reflects edit
r = srv.request('POST', '/api/v1/setup/samples/import?profile=qi_hitom_v1', [(FIX / 'valid-twenty.xlsx').read_bytes()],
                headers={'Content-Type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}).body
d, v = draft(r['samples'])
srv.json('POST', f"/api/v1/drafts/{d['draft_id']}/validate", dict(revision=d['revision']))
pv1 = srv.request('GET', f"/api/v1/drafts/{d['draft_id']}/preview?revision={d['revision']}")
cur = srv.json('GET', f"/api/v1/drafts/{d['draft_id']}").body; c = dict(cur['content']); s = [dict(x) for x in c['samples']]; s[0]['treatment'] = 'edited'; c['samples'] = s
d2 = srv.json('PATCH', f"/api/v1/drafts/{d['draft_id']}", dict(revision=cur['revision'], content=c)).body
stale = srv.request('GET', f"/api/v1/drafts/{d['draft_id']}/preview?revision={d['revision']}")
sub = srv.json('POST', '/api/v1/jobs', dict(draft_id=d['draft_id'], revision=d['revision'], plan_identity_sha256=pv1.body['plan_identity_sha256'], scientific_identity_sha256=pv1.body['scientific_identity_sha256']))
v2 = srv.json('POST', f"/api/v1/drafts/{d2['draft_id']}/validate", dict(revision=d2['revision'])).body
pv2 = srv.request('GET', f"/api/v1/drafts/{d['draft_id']}/preview?revision={d2['revision']}").body
R['edit_invalidates'] = dict(first_preview=pv1.status, stale_preview=(stale.status, (stale.body.get('error') or {}).get('code')),
                             stale_submit=(sub.status, (sub.body.get('error') or {}).get('code') if isinstance(sub.body, dict) else None),
                             revalidated=v2['status'], sci_changed=pv1.body['scientific_identity_sha256'] != pv2['scientific_identity_sha256'])
print('edit', R['edit_invalidates'], 'noexec', R['no_execution'])
# stale capability after restart, now with a fully valid draft
(OUT / 'restart').mkdir(); a = Adapter(OUT / 'restart'); z = a.upload('raw_R1.fastq.gz', [(fx / 'raw_R1.fastq.gz').read_bytes()]); z2 = a.upload('raw_R2.fastq.gz', [(fx / 'raw_R2.fastq.gz').read_bytes()]); stale_ids = (z.capability, z2.capability); a.close()
srv2 = Server(OUT / 'restart');
c = dict(tmpl); c.update(r1=stale_ids[0], r2=stale_ids[1]); dd = srv2.json('POST', '/api/v1/drafts').body
dd = srv2.json('PATCH', f"/api/v1/drafts/{dd['draft_id']}", dict(revision=dd['revision'], content=c)).body
vv = srv2.json('POST', f"/api/v1/drafts/{dd['draft_id']}/validate", dict(revision=dd['revision'])).body
R['restart_stale_capability'] = dict(status=vv['status'], issues=[(i['field'], i['code']) for i in vv['issues'] if i['field'] in ('r1', 'r2')])
print('restart', R['restart_stale_capability']); srv2.stop(); srv.stop()
(OUT / 'results.json').write_text(json.dumps(R, indent=1, default=str))
