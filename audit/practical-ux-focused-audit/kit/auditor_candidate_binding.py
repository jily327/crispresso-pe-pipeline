"""Independent auditor binding for the FROZEN audit/practical-ux/audit_candidate.py contract.

Binds the frozen adapter interface to candidate 04b9f88 transport:
  POST /api/v1/setup/uploads/fastq?filename=<name>  (raw application/octet-stream body)
  POST /api/v1/setup/samples/import?profile=...     (raw xlsx body; stateless parse)
A REAL uvicorn server runs in-process (thread) on loopback; requests use raw sockets so body
framing / disconnects are controlled by the auditor. Server-side fault barriers patch only
candidate module attributes inside this test process and only act on test-owned directories.
Nothing here implements candidate behaviour; inspection helpers read actual server state.
"""
import errno, gzip, hashlib, io, json, os, shutil, socket, sys, threading, time, tracemalloc, urllib.parse, uuid
from pathlib import Path
from types import SimpleNamespace
import pytest

HERE = Path(__file__).resolve().parent
PUX = HERE.parent
FROZEN = PUX / 'prep/audit/practical-ux'
FIX = FROZEN / 'fixtures'
TOKEN = 'audit-token'


def sha(b): return hashlib.sha256(b).hexdigest()


def stat_record(p):
    p = Path(p)
    if not os.path.lexists(p): return None
    s = os.lstat(p)
    rec = dict(ino=s.st_ino, mode=s.st_mode, mtime=s.st_mtime_ns, size=s.st_size, nlink=s.st_nlink)
    if os.path.isfile(p) and not os.path.islink(p): rec['sha256'] = sha(Path(p).read_bytes())
    return rec


class Server:
    def __init__(self, tmp, legacy_files=(), output_roots=(), runtime=None):
        import uvicorn
        from peanalyzer.gui.app import create_app
        self.upload_root = tmp / 'private' / 'uploads'
        self.app = create_app(token=TOKEN, setup_files=[str(p) for p in legacy_files], setup_output_roots=[str(p) for p in output_roots],
                              upload_root=self.upload_root, runtime=runtime)
        self.reads = []                       # ASGI http.request body sizes delivered to the application
        inner = self.app
        reads = self.reads
        async def wrapped(scope, receive, send):
            async def rec():
                m = await receive()
                if m.get('type') == 'http.request' and scope.get('path', '').endswith('/uploads/fastq'): reads.append(len(m.get('body', b'')))
                return m
            return await inner(scope, rec, send)
        s = socket.socket(); s.bind(('127.0.0.1', 0)); self.port = s.getsockname()[1]; s.close()
        self.server = uvicorn.Server(uvicorn.Config(wrapped, host='127.0.0.1', port=self.port, log_level='error', lifespan='on'))
        self.thread = threading.Thread(target=self.server.run, daemon=True); self.thread.start()
        end = time.monotonic() + 20
        while not self.server.started:
            assert time.monotonic() < end, 'server did not start'; time.sleep(.02)

    def stop(self):
        self.server.should_exit = True; self.thread.join(20)

    # -------- raw HTTP --------
    def request(self, method, path, body_chunks=(), headers=None, length='auto', close_after=None, host=None, origin=None, token=TOKEN):
        """Send a request; body_chunks iterable of bytes. close_after=N aborts after N body bytes (disconnect)."""
        chunks = list(body_chunks) if not hasattr(body_chunks, '__next__') else body_chunks
        h = {'Host': host or f'127.0.0.1:{self.port}', 'Connection': 'close'}
        if token is not None: h['X-Viewer-Token'] = token
        if origin: h['Origin'] = origin
        h.update(headers or {})
        if length == 'auto' and isinstance(chunks, list): h['Content-Length'] = str(sum(len(c) for c in chunks))
        elif isinstance(length, int): h['Content-Length'] = str(length)
        sock = socket.create_connection(('127.0.0.1', self.port), timeout=600)
        sock.sendall((f'{method} {path} HTTP/1.1\r\n' + ''.join(f'{k}: {v}\r\n' for k, v in h.items()) + '\r\n').encode())
        sent = 0
        try:
            for c in chunks:
                if close_after is not None and sent + len(c) > close_after:
                    c = c[:close_after - sent]
                    if c: sock.sendall(c); sent += len(c)
                    sock.shutdown(socket.SHUT_RDWR); sock.close(); return SimpleNamespace(status=None, body=None, aborted_after=sent)
                try: sock.sendall(c); sent += len(c)
                except (BrokenPipeError, ConnectionResetError): break
            if close_after is not None and sent >= close_after:
                sock.shutdown(socket.SHUT_RDWR); sock.close(); return SimpleNamespace(status=None, body=None, aborted_after=sent)
            data = b''
            while True:
                try: part = sock.recv(65536)
                except ConnectionResetError: break   # server answered/closed before draining the body
                if not part: break
                data += part
        finally:
            try: sock.close()
            except OSError: pass
        head, _, payload = data.partition(b'\r\n\r\n')
        status = int(head.split(b' ')[1]) if head else 'RESET'
        if b'transfer-encoding: chunked' in head.lower():
            out = b''; rest = payload
            while rest:
                n, _, rest = rest.partition(b'\r\n'); n = int(n, 16)
                if n == 0: break
                out += rest[:n]; rest = rest[n + 2:]
            payload = out
        try: body = json.loads(payload)
        except ValueError: body = payload
        return SimpleNamespace(status=status, body=body, headers=head.decode('latin1'))

    def json(self, method, path, value=None):
        b = json.dumps(value).encode() if value is not None else b''
        return self.request(method, path, [b] if b else [], headers={'Content-Type': 'application/json'} if b else {})


class Adapter:
    def __init__(self, tmp):
        self.tmp = Path(tmp); self.ext = self.tmp / 'external'; self.ext.mkdir()
        raw1 = (FIX / 'R1.fastq.gz').read_bytes(); raw2 = (FIX / 'R2.fastq.gz').read_bytes()
        self.raw1, self.raw2 = raw1, raw2
        # Explicitly seeded external victims (incl. the legacy launcher input, registered even if never selected).
        self.legacy = self.ext / 'legacy' / 'pool_R1.fastq.gz'; self.legacy.parent.mkdir(); self.legacy.write_bytes(raw1)
        self.other_analysis = self.ext / 'other-analysis' / 'R1.fastq.gz'; self.other_analysis.parent.mkdir(); self.other_analysis.write_bytes(raw1)
        self.workspace_file = self.ext / 'owned-workspace' / 'workflow' / 'plan.json'; self.workspace_file.parent.mkdir(parents=True); self.workspace_file.write_text('{"plan":1}\n')
        self.same_bytes_victim = self.ext / 'victim-same-bytes' / '.partial.gz'; self.same_bytes_victim.parent.mkdir(); self.same_bytes_victim.write_bytes(raw1)
        self.diff_victim = self.ext / 'victim-diff' / 'VICTIM.fastq.gz'; self.diff_victim.parent.mkdir(); self.diff_victim.write_bytes(raw2)
        self.out_root = self.tmp / 'outputs'; self.out_root.mkdir()
        self.server = Server(self.tmp, legacy_files=[self.legacy], output_roots=[self.out_root])
        self.root = self.server.upload_root; self.root.parent.mkdir(parents=True, exist_ok=True)
        self.victims = [self.legacy, self.other_analysis, self.workspace_file, self.same_bytes_victim, self.diff_victim]
        # Traversal landing spots relative to the upload storage / process cwd / filesystem root.
        for d in (self.root.parent, self.root.parent.parent, self.tmp, Path('/tmp')):
            v = d / 'VICTIM.fastq.gz'
            if not v.exists(): v.write_bytes(raw2)
            self.victims.append(v)
        self.full_hashes = {sha(raw1), sha(raw2)}
        self.refs = self.server.app.state.drafts.references
        self.log = []

    # ---------- inspection (private, read-only) ----------
    def victim_snapshot(self):
        snap = {str(v): stat_record(v) for v in self.victims}
        for d in {v.parent for v in self.victims} | {self.ext}:
            snap['dir:' + str(d)] = sorted(x for x in os.listdir(d) if not (d == self.root.parent and x == 'uploads'))
        return snap

    def accepted_capabilities(self):
        r = self.server.request('GET', '/api/v1/setup/capabilities')
        assert r.status == 200, r.body
        return sorted((f['file_id'], f['kind'], f.get('origin'), f.get('sha256'), f.get('size_bytes')) for f in r.body['files'])

    def accepted_partial_files(self):
        out = []
        for key, entry in self.refs.entries.items():
            if getattr(entry, 'origin', None) == 'upload':
                p = self.refs.paths[key]; b = p.read_bytes()
                if sha(b) not in self.full_hashes or entry.sha256 != sha(b): out.append((key, str(p), len(b)))
        return out

    def settle(self, timeout=20):
        end = time.monotonic() + timeout
        while self.server.app.state.uploads.lock.locked() and time.monotonic() < end: time.sleep(.05)
        time.sleep(.2)

    def staging_safe_after_settle(self):
        self.settle()
        registered = {str(p) for k, p in self.refs.paths.items() if getattr(self.refs.entries[k], 'origin', None) == 'upload'}
        if not self.root.exists(): return True
        for d in os.listdir(self.root):
            dp = self.root / d
            if os.path.islink(dp) or not dp.is_dir(): self.log.append(('unexpected-root-entry', d)); return False
            names = sorted(os.listdir(dp))
            if not any(str(dp / n) in registered for n in names): self.log.append(('leftover', d, names)); return False
            if any(n.startswith('.partial') for n in names): self.log.append(('partial-left', d, names)); return False
        return True

    def registered_bytes(self, cap): return self.refs.paths[cap].read_bytes()

    def registered_identity(self, cap):
        p = self.refs.paths[cap]; h = hashlib.sha256()
        with p.open('rb') as f:
            for b in iter(lambda: f.read(1 << 20), b''): h.update(b)
        return {'sha256': h.hexdigest(), 'size_bytes': p.stat().st_size}

    def is_owned_regular_file(self, cap):
        p = self.refs.paths[cap]; s = os.lstat(p)
        import stat as st
        parents_ok = all(not os.path.islink(x) for x in [p, *p.parents])
        return (st.S_ISREG(s.st_mode) and s.st_uid == os.getuid() and s.st_nlink == 1 and parents_ok
                and p.resolve().is_relative_to(self.root.resolve()) and st.S_IMODE(os.stat(self.root).st_mode) & 0o077 == 0)

    def path_controlled_by_client(self, cap, name):
        p = self.refs.paths[cap]; rel = p.resolve().relative_to(self.root.resolve()).as_posix()
        frags = {x for x in name.replace('\\', '/').split('/') if x and x not in ('.', '..')}
        return (not p.resolve().is_relative_to(self.root.resolve())) or any(f in rel for f in frags if f not in ('input.fastq.gz', 'input.fastq'))

    # ---------- transport ----------
    def _q(self, name): return '/api/v1/setup/uploads/fastq?filename=' + urllib.parse.quote(name, safe='')

    def upload(self, name, chunks, **kw):
        before = set(self.refs.entries)
        r = self.server.request('POST', self._q(name), list(chunks), headers={'Content-Type': 'application/octet-stream'}, **kw)
        new = set(self.refs.entries) - before
        cap = r.body.get('file_id') if isinstance(r.body, dict) else None
        accepted = r.status == 200 and cap in new
        self.log.append(('upload', name[:80], r.status, (r.body.get('error') or {}).get('code') if isinstance(r.body, dict) and r.body.get('error') else None))
        return SimpleNamespace(accepted=accepted, capability=cap if accepted else None, status=r.status, body=r.body)

    def upload_file_in_chunks(self, path, chunk_size):
        def gen():
            with open(path, 'rb') as f:
                for b in iter(lambda: f.read(chunk_size), b''): yield b
        size = Path(path).stat().st_size
        before = set(self.refs.entries); self.server.reads.clear()
        tracemalloc.start(); t0 = time.monotonic()
        r = self.server.request('POST', self._q(Path(path).name), gen(), headers={'Content-Type': 'application/octet-stream'}, length=size)
        cur, peak = tracemalloc.get_traced_memory(); tracemalloc.stop()
        self._stream = dict(size=size, chunk=chunk_size, peak_traced=peak, seconds=round(time.monotonic() - t0, 1), reads_n=len(self.server.reads),
                            reads_max=max(self.server.reads or [0]), rss_kb=int(open('/proc/self/status').read().split('VmHWM:')[1].split()[0]))
        self.log.append(('stream', self._stream))
        new = set(self.refs.entries) - before; cap = r.body.get('file_id') if isinstance(r.body, dict) else None
        return SimpleNamespace(accepted=r.status == 200 and cap in new, capability=cap, status=r.status, body=r.body)

    def interrupt_upload(self, path, boundary):
        raw = Path(path).read_bytes(); n = len(raw)
        chunks = [raw[i:i + 64] for i in range(0, n, 64)]
        if boundary == 'written-before-register':
            real = self.refs.register_upload
            def boom(*a, **k): raise RuntimeError('AUDIT injected registration failure after write')
            self.refs.register_upload = boom
            try: r = self.server.request('POST', self._q('R1.fastq.gz'), [raw], headers={'Content-Type': 'application/octet-stream'})
            finally: self.refs.register_upload = real
            self.log.append(('interrupt', boundary, r.status)); self.settle(); return r
        cut = {'zero': 0, 'first': 64, 'middle': n // 2, 'last-before-finalize': n - 1}[boundary]
        r = self.server.request('POST', self._q('R1.fastq.gz'), chunks, headers={'Content-Type': 'application/octet-stream'}, length=n, close_after=cut)
        self.log.append(('interrupt', boundary, r)); self.settle(); return r

    def upload_with_owned_staging_fault(self, attack, path):
        import peanalyzer.gui.uploads as U
        raw = Path(path).read_bytes(); real = U.validate_fastq; done = {}
        def barrier(p):
            p = Path(p); d = p.parent; key = d.name
            if attack == 'staging-symlink':        # partial -> symlink to external file with IDENTICAL bytes
                os.unlink(p); os.symlink(self.same_bytes_victim, p)
            elif attack == 'replace-temp':         # new inode, different (valid FASTQ) content
                tmpf = d / '.audit-new'; tmpf.write_bytes(self.raw2); os.replace(tmpf, p)
            elif attack == 'external-symlink':     # partial -> symlink to external victim (different bytes)
                os.unlink(p); os.symlink(self.diff_victim, p)
            elif attack == 'directory-substitution':  # owned <uuid> dir -> symlink to external dir holding identical bytes
                os.rename(d, d.with_name(key + '.audit-moved')); os.symlink(self.same_bytes_victim.parent, d)
            done['attack'] = attack
            return real(p)
        U.validate_fastq = barrier
        try: r = self.server.request('POST', self._q('R1.fastq.gz'), [raw], headers={'Content-Type': 'application/octet-stream'})
        finally: U.validate_fastq = real
        self.settle(); self.log.append(('staging-fault', attack, done, r.status, r.body if not isinstance(r.body, bytes) else None))
        # Auditor-owned cleanup of auditor-created artefacts only (never candidate/victim files).
        for e in list(os.listdir(self.root)):
            ep = self.root / e
            if os.path.islink(ep): os.unlink(ep)
        return r

    def send_path_instead_of_file(self, target):
        path = {'/etc/passwd': '/etc/passwd', 'owned-workspace': str(self.workspace_file), 'other-analysis': str(self.other_analysis), 'legacy-input': str(self.legacy)}[target]
        before = set(self.refs.entries); results = []
        b = path.encode()
        results.append(self.server.request('POST', self._q('x.fastq'), [b], headers={'Content-Type': 'application/octet-stream'}).status)
        results.append(self.server.request('POST', self._q(path), [b''], headers={'Content-Type': 'application/octet-stream'}).status)
        results.append(self.server.request('POST', self._q(path + '.fastq.gz'), [b], headers={'Content-Type': 'application/octet-stream'}).status)
        results.append(self.server.request('POST', self._q('x.fastq.gz'), [json.dumps({'path': path}).encode()], headers={'Content-Type': 'application/json'}).status)
        bd = 'auditboundary'
        mp = (f'--{bd}\r\nContent-Disposition: form-data; name="file"; filename="{path}"\r\nContent-Type: application/gzip\r\n\r\n{path}\r\n--{bd}--\r\n').encode()
        results.append(self.server.request('POST', self._q('x.fastq.gz'), [mp], headers={'Content-Type': f'multipart/form-data; boundary={bd}'}).status)
        results.append(self.server.request('POST', self._q('x.fastq'), [path.encode()], headers={'Content-Type': 'text/plain'}).status)
        # Draft reference fields must not dereference a server path either.
        d = self.server.json('POST', '/api/v1/drafts').body
        p = self.server.json('PATCH', f"/api/v1/drafts/{d['draft_id']}", dict(revision=d['revision'], content=dict(r1=path, r2=path))).body
        draft_ok = isinstance(p, dict) and p.get('content', {}).get('r1') == path
        if draft_ok:
            v = self.server.json('POST', f"/api/v1/drafts/{d['draft_id']}/validate", dict(revision=p['revision'])).body
            draft_ok = not any(i.get('field') == 'r1' and i.get('severity') == 'ERROR' for i in v.get('issues', []))
        new = set(self.refs.entries) - before
        self.log.append(('path-injection', target, results, draft_ok))
        return SimpleNamespace(accepted=bool(new) or draft_ok, statuses=results)

    def register_external_fixture(self):
        k = next(k for k, p in self.refs.paths.items() if p == self.legacy); return k

    def replace_selected_upload(self, first, second):
        d = self.server.json('POST', '/api/v1/drafts').body; legacy = self.register_external_fixture()
        p = self.server.json('PATCH', f"/api/v1/drafts/{d['draft_id']}", dict(revision=d['revision'], content=dict(r1=first, r2=legacy))).body
        p = self.server.json('PATCH', f"/api/v1/drafts/{d['draft_id']}", dict(revision=p['revision'], content=dict(r1=second, r2=legacy))).body
        self.server.json('DELETE', f"/api/v1/drafts/{d['draft_id']}")

    def cleanup_uploads(self):
        self.server.stop(); self._stopped = True

    def external_exists(self, cap): return self.legacy.is_file() and self.legacy.read_bytes() == self.raw1

    # ---------- Excel ----------
    def draft_snapshot(self):
        if not hasattr(self, 'draft'):
            d = self.server.json('POST', '/api/v1/drafts').body
            seed = [dict(sample_name='seed', barcode_id='1A', forward_barcode='', reverse_barcode='', editor='e', treatment='t', dpi='6', biological_replicate='01', paired=False, block_id='')]
            self.draft = self.server.json('PATCH', f"/api/v1/drafts/{d['draft_id']}", dict(revision=d['revision'], content=dict(samples=seed))).body
        v = self.server.json('GET', f"/api/v1/drafts/{self.draft['draft_id']}").body
        return json.dumps(v, sort_keys=True)

    def import_workbook(self, path, profile='qi_hitom_v1'):
        r = self.server.request('POST', f'/api/v1/setup/samples/import?profile={profile}', [Path(path).read_bytes()],
                                headers={'Content-Type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'})
        body = r.body if isinstance(r.body, dict) else {}
        accepted = r.status == 200 and not body.get('errors') and bool(body.get('samples'))
        self._imported = body.get('samples') if accepted else None
        self.log.append(('import', Path(path).name, profile, r.status, [(e['row'], e['field'], e['message'][:80]) for e in body.get('errors', [])][:4]))
        if accepted:   # mirror the candidate client contract: apply only an error-free import to the draft
            d = json.loads(self.draft_snapshot())
            self.draft = self.server.json('PATCH', f"/api/v1/drafts/{d['draft_id']}", dict(revision=d['revision'], content=dict(samples=body['samples']))).body
        return SimpleNamespace(accepted=accepted, status=r.status, body=body)

    @staticmethod
    def _qualified(rows, profile='qi_hitom_v1'):
        from peanalyzer.validation import load_samples
        from peanalyzer.gui.setup_validation import MemoryCSV
        from peanalyzer.barcodes import load_catalog, bundled_catalog, Catalog
        cat = load_catalog(bundled_catalog()) if profile == 'qi_hitom_v1' else Catalog([], authoritative=False)
        return sorted([s.model_dump() for s in load_samples(MemoryCSV(rows, 'audit'), cat)], key=lambda s: s['barcode_id'])

    def canonical_imported_samples(self):
        return self._qualified(self._imported)

    def authoritative_manual_equivalent(self, case):
        """Frozen cell values -> values a person would type into the existing manual Sample form.
        Auditor reads the xlsx itself (openpyxl) and resolves catalog-only rows through the EXISTING
        catalog authority (Catalog.resolve_sample with end IDs), never through the parser under test."""
        from openpyxl import load_workbook
        from peanalyzer.barcodes import load_catalog, bundled_catalog
        from peanalyzer.schemas import Sample
        b = load_workbook(FIX / (case + '.xlsx')); s = b['Samples']; rows = list(s.iter_rows(values_only=True))
        hdr = [h.strip() for h in rows[0]]; out = []
        for r in rows[1:]:
            v = {k: ('' if x is None else str(x).strip()) for k, x in zip(hdr, r)}
            if not any(v.values()): continue
            if not v['forward_barcode'] and not v['reverse_barcode']:
                c = load_catalog(bundled_catalog()); bid = v['barcode_id']
                s_ = c.resolve_sample(Sample(sample_name=v['sample_name'], barcode_id=bid, forward_barcode_id='F' + bid[:-1], reverse_barcode_id='R' + bid[-1],
                                             editor=v['editor'], treatment=v['treatment'], dpi=int(v['dpi']), biological_replicate=v['biological_replicate'], paired=False))
                v['forward_barcode'], v['reverse_barcode'] = s_.forward_barcode, s_.reverse_barcode
            v['paired'] = v['paired'].lower() == 'true'
            out.append({k: v.get(k, '') for k in ('sample_name', 'barcode_id', 'forward_barcode', 'reverse_barcode', 'editor', 'treatment', 'dpi', 'biological_replicate', 'paired', 'block_id')})
        return self._qualified(out)

    def import_with_catalog_probe(self, path, end_id, value):
        import peanalyzer.gui.sample_workbook as W
        from peanalyzer.barcodes import Catalog, bundled_catalog
        real_load = W.load_catalog; real_resolve = Catalog.resolve_sample; calls = dict(load=0, resolve=0)
        before = sha(Path(bundled_catalog()).read_bytes())
        def load(p):
            calls['load'] += 1; c = real_load(p)
            if c.authoritative: c.ends['forward'][end_id] = value     # fresh in-memory instance only
            return c
        def resolve(self_, sample): calls['resolve'] += 1; return real_resolve(self_, sample)
        W.load_catalog = load; Catalog.resolve_sample = resolve
        try:
            r = self.server.request('POST', '/api/v1/setup/samples/import?profile=qi_hitom_v1', [Path(path).read_bytes()],
                                    headers={'Content-Type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'})
        finally: W.load_catalog = real_load; Catalog.resolve_sample = real_resolve
        self.probe = dict(calls=calls, catalog_file_unchanged=sha(Path(bundled_catalog()).read_bytes()) == before, errors=r.body.get('errors'))
        self.log.append(('catalog-probe', self.probe))
        return r.body['samples'][0]['forward_barcode'] if r.body.get('samples') else None

    # ---------- streaming ----------
    def generate_large_fastq(self, tmp_path, minimum_compressed_bytes):
        sys.path.insert(0, str(FROZEN)); import generate_fixtures as G
        pairs = 1_000_000
        while True:
            out = Path(tmp_path) / f'large-{pairs}'; G.generate(out, pairs)
            p = out / 'R1.fastq.gz'
            if p.stat().st_size >= minimum_compressed_bytes: self.log.append(('large', pairs, p.stat().st_size)); return p
            pairs = int(pairs * 1.5)

    def streaming_evidence(self):
        s = self._stream
        bounded = s['reads_max'] <= 1024 * 1024 and s['reads_n'] > 1 and s['peak_traced'] < 32 * 1024 * 1024
        return SimpleNamespace(bounded_reads_and_no_linear_rss_growth=bounded, **s)

    # ---------- not bindable in-process (bound separately; explicit errors, never passes) ----------
    def __getattr__(self, name):
        if name.startswith('_') or name in ('draft', 'probe'): raise AttributeError(name)
        raise NotImplementedError(f'AUDIT BINDING: {name} is bound in the separate browser/execution harness (see audit report), not here')

    def close(self):
        if not getattr(self, '_stopped', False): self.server.stop()


@pytest.fixture
def audit_adapter(tmp_path):
    a = Adapter(tmp_path)
    yield a
    out = os.environ.get('AUDIT_LOG_DIR')
    if out:
        Path(out).mkdir(parents=True, exist_ok=True)
        with open(Path(out) / 'adapter-log.jsonl', 'a') as f:
            f.write(json.dumps(dict(test=os.environ.get('PYTEST_CURRENT_TEST'), log=a.log), default=str) + '\n')
    a.close()
